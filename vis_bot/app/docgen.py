"""Сборка Word-файла (.docx) для заметки о личности."""
from __future__ import annotations

import re
import tempfile

from docx import Document

# Убираем HTML-теги, если оператор вставил их в текст (в .docx они не нужны).
_TAGS = re.compile(r"</?[^>]+>")


def _clean(text: str) -> str:
    return _TAGS.sub("", text or "").strip()


def build_person_docx(entry: dict) -> tuple[str, str]:
    """Создаёт .docx по записи личности. Возвращает (путь_к_файлу, имя_файла)."""
    name = _clean(entry.get("name", "Личность"))
    years = _clean(entry.get("years", ""))

    doc = Document()
    doc.add_heading(name, level=0)
    if years:
        doc.add_paragraph(years)

    bio = _clean(entry.get("bio", ""))
    if bio:
        doc.add_paragraph(bio)

    achievements = entry.get("achievements") or []
    if achievements:
        doc.add_heading("Достижения", level=1)
        for item in achievements:
            doc.add_paragraph(_clean(item), style="List Bullet")

    quotes = entry.get("quotes") or []
    if quotes:
        doc.add_heading("Цитаты", level=1)
        for quote in quotes:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(f"«{_clean(quote)}»")
            run.italic = True

    fd, path = tempfile.mkstemp(suffix=".docx")
    import os

    os.close(fd)
    doc.save(path)

    safe_name = re.sub(r"[^\w\-. ]", "_", name).strip() or "person"
    return path, f"{safe_name}.docx"
