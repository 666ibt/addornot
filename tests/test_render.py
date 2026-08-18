"""Formula generation, block cloning and the fixes applied to the template."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from compsheet.config import Config, SpecItem
from compsheet.errors import FxError
from compsheet.fx import FxTable
from compsheet.models import Field, LineItem, Offer
from compsheet.template.render import render


def make_offer(name, currency="USD", vat_included=False, prices=(100.0, 200.0)):
    offer = Offer(
        company_name=Field.extracted(name),
        currency=Field.extracted(currency),
        vat_included=Field.extracted(vat_included),
        incoterms=Field.extracted("DDP"),
        lead_time=Field.assumed("30 дней", "выведено из формулировки «месяц»"),
    )
    offer.source_files = [f"/offers/{name}.pdf"]
    offer.assumptions = ["lead_time: выведено из формулировки «месяц»"]
    for index, price in enumerate(prices):
        offer.line_items.append(
            LineItem(
                description=Field.extracted(f"позиция {index + 1}"),
                qty=Field.extracted(2.0),
                unit_price=Field.extracted(price),
                total_price=Field.extracted(price * 2),
                spec_index=index,
            )
        )
    return offer


@pytest.fixture
def three_position_config():
    config = Config.load("fixtures/config.yaml")
    config.spec = [
        SpecItem("Позиция А", 4, "шт"),
        SpecItem("Позиция Б", 2, "шт"),
        SpecItem("Позиция В", 6, "компл"),
    ]
    return config


class TestBlockLayout:
    def test_offers_are_grouped_two_per_block(self, template_path, config, fx, tmp_path):
        offers = [make_offer(f"Компания {n}") for n in range(5)]
        out = tmp_path / "out.xlsx"
        report = render(str(template_path), offers, config, fx, str(out))
        assert report.blocks == 3  # 2 + 2 + 1
        assert len(report.placements) == 5

    def test_extra_positions_push_the_totals_down(self, template_path, three_position_config,
                                                  fx, tmp_path):
        """The template has one nomenclature row; three positions must fit."""
        offers = [make_offer("Компания", prices=(100.0, 200.0, 300.0))]
        out = tmp_path / "out.xlsx"
        report = render(str(template_path), offers, three_position_config, fx, str(out))
        placement = report.placements[0]
        assert len(placement.item_rows) == 3
        # Every totals row sits below the last nomenclature row.
        assert min(placement.role_rows[role] for role in
                   ("total_qty", "total_amount_ccy", "grand_total")) > max(placement.item_rows)

    def test_labels_and_styles_survive_cloning(self, template_path, config, fx, tmp_path):
        out = tmp_path / "out.xlsx"
        render(str(template_path), [make_offer("А"), make_offer("Б"), make_offer("В")],
               config, fx, str(out))
        sheet = load_workbook(out)["Конкурентный лист"]
        labels = [sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)]
        assert labels.count("Принадлежность поставщика") == 2  # one per block
        assert sheet.cell(row=4, column=5).font.bold


class TestFormulas:
    @pytest.fixture
    def sheet(self, template_path, three_position_config, fx, tmp_path):
        out = tmp_path / "out.xlsx"
        self.report = render(str(template_path), [make_offer("Компания", prices=(1.0, 2.0, 3.0))],
                             three_position_config, fx, str(out))
        return load_workbook(out)["Конкурентный лист"]

    def test_totals_sum_the_whole_range_not_a_single_row(self, sheet):
        placement = self.report.placements[0]
        first, last = placement.item_rows[0], placement.item_rows[-1]
        formula = sheet.cell(row=placement.role_rows["total_amount_ccy"], column=5).value
        assert formula == f"=SUM(F{first}:F{last})"

    def test_amount_is_quantity_times_price(self, sheet):
        row = self.report.placements[0].item_rows[0]
        assert sheet.cell(row=row, column=6).value == f"=D{row}*E{row}"

    def test_local_currency_uses_the_rate_table(self, sheet):
        placement = self.report.placements[0]
        formula = sheet.cell(row=placement.role_rows["total_amount_base"], column=5).value
        assert "VLOOKUP" in formula
        assert f"E{placement.role_rows['currency']}" in formula

    def test_vat_is_added_only_when_the_price_excludes_it(self, sheet):
        placement = self.report.placements[0]
        formula = sheet.cell(row=placement.role_rows["vat_amount"], column=5).value
        assert formula.startswith("=IF(") and 'без НДС' in formula and "0.12" in formula

    def test_only_libreoffice_safe_functions_are_used(self, sheet):
        banned = ("XLOOKUP", "FILTER(", "UNIQUE(", "LET(", "TEXTJOIN")
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert not any(name in cell.value.upper() for name in banned)


class TestFixesAppliedToTheTemplate:
    @pytest.fixture
    def book(self, template_path, config, fx, tmp_path):
        out = tmp_path / "out.xlsx"
        self.report = render(str(template_path), [make_offer("А"), make_offer("Б")],
                             config, fx, str(out))
        return load_workbook(out)

    def test_rate_table_is_filled_in(self, book):
        reference = book["Справочники"]
        codes = [reference.cell(row=row, column=6).value for row in range(2, 8)]
        assert "USD" in codes and "UZS" in codes
        assert reference.cell(row=codes.index("USD") + 2, column=7).value == 12950.0

    def test_fit_to_one_page_is_switched_off(self, book):
        sheet = book["Конкурентный лист"]
        assert sheet.page_setup.fitToHeight == 0
        assert sheet.page_setup.fitToWidth == 1
        assert any("одну страницу" in note for note in self.report.notes)

    def test_the_company_cell_carries_the_audit_trail(self, book):
        comment = book["Конкурентный лист"].cell(row=4, column=5).comment
        assert comment is not None
        assert "Файлы" in comment.text and "Допущения" in comment.text

    def test_an_assumed_field_is_marked_in_its_own_cell(self, book):
        row = self.report.placements[0].role_rows["lead_time"]
        comment = book["Конкурентный лист"].cell(row=row, column=5).comment
        assert comment is not None and comment.text.startswith("ПРЕДПОЛОЖЕНО")


class TestGuards:
    def test_a_currency_without_a_rate_stops_the_run(self, template_path, config, tmp_path):
        table = FxTable(base="UZS", rates={"USD": 12950})
        with pytest.raises(FxError, match="TRY"):
            render(str(template_path), [make_offer("Т", currency="TRY")], config, table,
                   str(tmp_path / "out.xlsx"))

    def test_a_position_nobody_quoted_is_written_explicitly(self, template_path,
                                                            three_position_config, fx, tmp_path):
        from compsheet.pipeline import fill_missing_from_spec

        offer = make_offer("Компания", prices=(100.0,))
        fill_missing_from_spec(offer, three_position_config)
        out = tmp_path / "out.xlsx"
        report = render(str(template_path), [offer], three_position_config, fx, str(out))
        sheet = load_workbook(out)["Конкурентный лист"]
        rows = report.placements[0].item_rows
        assert sheet.cell(row=rows[1], column=5).value == 0
        assert sheet.cell(row=rows[1], column=5).comment is not None
