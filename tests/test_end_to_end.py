"""Golden run: the fixture bundle, start to finish, checked after recalculation.

These assertions are the regression contract -- extracted values, generated
formulas and LibreOffice's own arithmetic all have to agree with the numbers
printed in the fixture offers.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from compsheet.template.render import render
from compsheet.validate import recalculate, scan_formula_errors
from conftest import needs_libreoffice, needs_ocr

# Company -> (total in its own currency, currency, VAT included in the price)
EXPECTED = {
    "Альфа": (136400.0, "USD", False),
    "Эпсилон": (139200.0, "USD", False),
    "GAMMA": (67600.0, "EUR", None),
}
RATES = {"USD": 12950.0, "EUR": 14100.0}
VAT_RATE = 0.12


def find_by_file(result, fragment):
    """Some suppliers only identify themselves through their filenames."""
    matches = [offer for offer in result.offers
               if any(fragment in path for path in offer.source_files)]
    assert matches, f"no offer with a file matching {fragment!r}"
    return matches[0]


def find_offer(result, fragment):
    matches = [offer for offer in result.offers if fragment.lower() in offer.display_name.lower()]
    assert matches, f"no offer matching {fragment!r}"
    return matches[0]


@needs_ocr
class TestExtraction:
    def test_every_supplier_bundle_becomes_one_offer(self, extracted):
        assert len(extracted.offers) == 5

    def test_a_truncated_download_is_reported_not_skipped(self, extracted):
        assert [f.path for f in extracted.failed_files if "zeta" in f.path]

    def test_two_files_from_one_supplier_are_merged(self, extracted):
        offer = find_offer(extracted, "Эпсилон")
        assert len(offer.source_files) == 2

    def test_a_mangled_docx_is_rerouted_through_ocr(self, extracted):
        offer = find_by_file(extracted, "delta")
        assert any("mangled" in w or "OCR" in w for w in offer.warnings)

    @pytest.mark.parametrize("fragment", list(EXPECTED))
    def test_prices_match_the_documents(self, extracted, fragment):
        offer = find_offer(extracted, fragment)
        total, currency, vat = EXPECTED[fragment]
        assert offer.currency.value == currency
        assert offer.vat_included.value is vat
        assert offer.stated_total.value == pytest.approx(total)

    def test_a_position_nobody_quoted_is_flagged_and_zero_filled(self, extracted):
        offer = find_offer(extracted, "GAMMA")
        second = next(item for item in offer.line_items if item.spec_index == 1)
        assert second.qty.value == 0
        assert any("не предложена" in w for w in offer.warnings)

    def test_a_short_quantity_is_flagged(self, extracted):
        """Delta priced one unit where the specification asks for four."""
        offer = find_by_file(extracted, "delta")
        assert any("требуется 4" in w for w in offer.warnings)

    def test_an_unreadable_table_is_left_empty_rather_than_guessed(self, extracted):
        offer = find_offer(extracted, "Бета")
        assert all(item.qty.value is None for item in offer.line_items)


@pytest.fixture(scope="module")
def built(extracted, config, fx, template_path, workdir):
    """Render the fixture set and let LibreOffice compute every formula."""
    out = workdir / "golden.xlsx"
    report = render(str(template_path), extracted.offers, config, fx, str(out))
    recalculated = recalculate(out, workdir / "golden-recalc")
    sheet = load_workbook(recalculated, data_only=True)["Конкурентный лист"]
    return report, sheet, recalculated


@needs_libreoffice
@needs_ocr
class TestGeneratedWorkbook:
    def test_no_formula_errors_survive_recalculation(self, built):
        assert scan_formula_errors(built[2]) == []

    def test_block_count_follows_suppliers_per_block(self, built, extracted, config):
        report = built[0]
        expected = -(-len(extracted.offers) // config.suppliers_per_block)
        assert report.blocks == expected

    @pytest.mark.parametrize("fragment", list(EXPECTED))
    def test_totals_agree_with_the_supplier_document(self, built, fragment):
        report, sheet, _ = built
        total, currency, _ = EXPECTED[fragment]
        placement = next(p for p in report.placements if fragment.lower() in p.company.lower())
        computed = sheet.cell(row=placement.role_rows["total_amount_ccy"],
                              column=placement.column).value
        assert computed == pytest.approx(total)

    @pytest.mark.parametrize("fragment", list(EXPECTED))
    def test_conversion_to_the_base_currency_uses_the_real_rate(self, built, fragment):
        report, sheet, _ = built
        total, currency, _ = EXPECTED[fragment]
        placement = next(p for p in report.placements if fragment.lower() in p.company.lower())
        computed = sheet.cell(row=placement.role_rows["total_amount_base"],
                              column=placement.column).value
        assert computed == pytest.approx(total * RATES[currency])

    def test_vat_is_added_only_for_prices_quoted_without_it(self, built):
        report, sheet, _ = built
        for fragment, (total, currency, vat_included) in EXPECTED.items():
            placement = next(p for p in report.placements
                             if fragment.lower() in p.company.lower())
            vat = sheet.cell(row=placement.role_rows["vat_amount"],
                             column=placement.column).value
            expected = total * RATES[currency] * VAT_RATE if vat_included is False else 0
            assert vat == pytest.approx(expected), fragment

    def test_grand_total_adds_base_vat_and_transport(self, built):
        report, sheet, _ = built
        placement = next(p for p in report.placements if "Альфа" in p.company)
        parts = [sheet.cell(row=placement.role_rows[role], column=placement.column).value
                 for role in ("total_amount_base", "vat_amount", "transport_cost")]
        grand = sheet.cell(row=placement.role_rows["grand_total"], column=placement.column).value
        assert grand == pytest.approx(sum(parts))

    def test_quantities_sum_over_every_position(self, built):
        report, sheet, _ = built
        placement = next(p for p in report.placements if "Альфа" in p.company)
        assert sheet.cell(row=placement.role_rows["total_qty"],
                          column=placement.column).value == 6  # 4 + 2


@needs_libreoffice
@needs_ocr
class TestCorrectionRoundTrip:
    def test_a_single_field_can_be_fixed_without_re_extracting(self, extracted, config, fx,
                                                               template_path, workdir, tmp_path):
        """The unparsed OCR row is corrected through the review JSON."""
        from compsheet.models import ExtractionResult
        from compsheet.review import apply_edit, load_json, save_json

        path = tmp_path / "offers.json"
        save_json(extracted, path)
        reloaded = load_json(path)
        assert isinstance(reloaded, ExtractionResult)

        for edit in ("Бета.item1.qty=4", "Бета.item1.unit_price=235000000",
                     "Бета.item2.qty=2", "Бета.item2.unit_price=398000000"):
            apply_edit(reloaded, edit)

        out = tmp_path / "fixed.xlsx"
        report = render(str(template_path), reloaded.offers, config, fx, str(out))
        sheet = load_workbook(recalculate(out, tmp_path), data_only=True)["Конкурентный лист"]
        placement = next(p for p in report.placements if "Бета" in p.company)
        computed = sheet.cell(row=placement.role_rows["total_amount_ccy"],
                              column=placement.column).value
        assert computed == pytest.approx(1_736_000_000.0)
