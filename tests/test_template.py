"""The template must be *read*, not assumed."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from compsheet.config import Config
from compsheet.errors import TemplateError
from compsheet.template.parse import ITEM, PARAM, TOTAL, parse_block


@pytest.fixture(scope="module")
def block(template_path, config):
    sheet = load_workbook(template_path)["Конкурентный лист"]
    return parse_block(sheet, config)


class TestParsedGeometry:
    def test_label_column_and_supplier_spans_come_from_the_merges(self, block):
        assert block.label_column == 1
        assert block.supplier_columns == [(5, 6), (7, 8)]

    def test_parameter_rows_are_found_by_role_not_by_count(self, block):
        roles = [row.role for row in block.rows if row.kind == PARAM]
        assert "currency" in roles and "incoterms" in roles and "vat_included" in roles
        assert len(roles) == len(set(roles))

    def test_totals_are_identified_including_the_grand_total(self, block):
        roles = [row.role for row in block.rows if row.kind == TOTAL]
        for expected in ("total_qty", "total_amount_ccy", "total_amount_base",
                         "vat_amount", "transport_cost", "grand_total"):
            assert expected in roles

    def test_the_single_sample_item_row_is_the_repeat_unit(self, block):
        assert len([row for row in block.rows if row.kind == ITEM]) == 1

    def test_shared_and_per_supplier_item_columns_are_separated(self, block):
        assert block.item_columns["qty"] == 4
        assert block.supplier_item_columns[0]["amount"] == 6
        assert block.supplier_item_columns[1]["amount"] == 8

    def test_row_heights_and_styles_are_captured(self, block):
        header = block.rows[0]
        assert header.height == 30
        assert all(cell.style is not None for cell in header.cells)


class TestParseFailures:
    def test_a_sheet_without_labels_is_rejected(self, template_path, config):
        sheet = load_workbook(template_path)["Справочники"]
        with pytest.raises(TemplateError):
            parse_block(sheet, config)

    def test_custom_labels_can_be_taught_through_the_config(self, template_path):
        config = Config.from_dict({"label_patterns": {"currency": [r"валюта ткп"]}})
        sheet = load_workbook(template_path)["Конкурентный лист"]
        assert parse_block(sheet, config).role_row("currency") is not None
