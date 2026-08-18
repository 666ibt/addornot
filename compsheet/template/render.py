"""Write the comparison blocks into a copy of the template workbook.

The template supplies geometry and styling only.  Every formula is generated
from the row's semantic role against the actual row numbers of the block being
written, which is what makes a variable number of nomenclature positions -- and
a working local-currency conversion -- possible.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from ..config import Config
from ..errors import FxError, TemplateError
from ..fx import FxTable
from ..models import ASSUMED, MISSING, Offer
from .parse import HEADER, ITEM, ITEM_HEADER, OTHER, PARAM, TOTAL, BlockTemplate, parse_block

FX_RANGE_NAME = "compsheet_fx"
NOT_STATED = "— не указано —"
VAT_IN_PRICE = "с НДС"
VAT_NOT_IN_PRICE = "без НДС"


@dataclass
class Placement:
    """Where one offer ended up, so validation can find its totals again."""

    company: str
    sheet: str
    column: int
    role_rows: dict[str, int] = field(default_factory=dict)
    item_rows: list[int] = field(default_factory=list)


@dataclass
class RenderReport:
    blocks: int = 0
    rows_written: int = 0
    item_rows_per_block: int = 0
    fx_rows: int = 0
    notes: list[str] = field(default_factory=list)
    placements: list[Placement] = field(default_factory=list)


def _letter(column: int) -> str:
    return get_column_letter(column)


def find_main_sheet(workbook, config: Config):
    """The sheet whose labels match the configured block roles."""
    errors = []
    for sheet in workbook.worksheets:
        try:
            return sheet, parse_block(sheet, config)
        except TemplateError as error:
            errors.append(f"{sheet.title}: {error}")
    raise TemplateError(
        "no sheet in the template looks like a comparison block:\n  " + "\n  ".join(errors)
    )


def write_fx_table(workbook, fx: FxTable, config: Config) -> tuple[str, int]:
    """Write the rate table used by the conversion formulas; return its range."""
    rows = fx.as_rows()
    target = None
    anchor = None
    if config.reference_sheet and config.reference_sheet in workbook.sheetnames:
        target = workbook[config.reference_sheet]

    candidates = [target] if target is not None else list(workbook.worksheets)
    for sheet in candidates:
        for row in range(1, min(sheet.max_row, 50) + 1):
            for column in range(1, min(sheet.max_column, 40) + 1):
                value = sheet.cell(row=row, column=column).value
                if isinstance(value, str) and any(
                    key in value.lower() for key in ("курс", "rate", "kurs")
                ):
                    target, anchor = sheet, (row, column)
                    break
            if anchor:
                break
        if anchor:
            break

    if anchor is None:
        target = target or (workbook.worksheets[1] if len(workbook.worksheets) > 1
                            else workbook.create_sheet("Курсы валют"))
        column = (target.max_column or 0) + 2
        anchor = (1, column + 1)
        target.cell(row=1, column=column, value="Валюта")
        target.cell(row=1, column=column + 1, value=f"Курс к {fx.base}")

    header_row, rate_column = anchor
    code_column = max(1, rate_column - 1)
    for offset, (code, rate) in enumerate(rows, start=1):
        target.cell(row=header_row + offset, column=code_column, value=code)
        cell = target.cell(row=header_row + offset, column=rate_column, value=rate)
        cell.number_format = "#,##0.0000"
    # Clear stale rows left over from a previous, longer table.
    extra = header_row + len(rows) + 1
    while target.cell(row=extra, column=code_column).value is not None:
        target.cell(row=extra, column=code_column).value = None
        target.cell(row=extra, column=rate_column).value = None
        extra += 1

    note = target.cell(row=header_row + len(rows) + 2, column=code_column)
    note.value = f"Источник курсов: {fx.source}{(' от ' + fx.date) if fx.date else ''}"

    reference = (
        f"{quote_sheetname(target.title)}!"
        f"${_letter(code_column)}${header_row + 1}:"
        f"${_letter(rate_column)}${header_row + len(rows)}"
    )
    if FX_RANGE_NAME in workbook.defined_names:
        del workbook.defined_names[FX_RANGE_NAME]
    workbook.defined_names.add(DefinedName(FX_RANGE_NAME, attr_text=reference))
    return reference, len(rows)


def _capture_validations(sheet, block: BlockTemplate):
    """Dropdowns that live inside the template block, as (formula, row offsets)."""
    captured = []
    for validation in list(sheet.data_validations.dataValidation):
        offsets = []
        for cell_range in validation.sqref.ranges:
            if cell_range.min_row >= block.top and cell_range.max_row <= block.bottom:
                offsets.append((
                    cell_range.min_row - block.top, cell_range.min_col,
                    cell_range.max_row - block.top, cell_range.max_col,
                ))
        if offsets:
            captured.append((validation.type, validation.formula1, offsets))
    return captured


def _clear_from(sheet, block: BlockTemplate) -> None:
    """Remove the sample block and anything after it."""
    for merged in list(sheet.merged_cells.ranges):
        if merged.max_row >= block.top:
            sheet.unmerge_cells(str(merged))
    sheet.data_validations.dataValidation = [
        validation for validation in sheet.data_validations.dataValidation
        if all(cell_range.max_row < block.top for cell_range in validation.sqref.ranges)
    ]
    if sheet.max_row >= block.top:
        sheet.delete_rows(block.top, sheet.max_row - block.top + 1)
    for row in list(sheet.row_dimensions):
        if row >= block.top:
            del sheet.row_dimensions[row]


def _field_text(value, missing_text: str = NOT_STATED) -> Any:
    if value is None or value == "":
        return missing_text
    if isinstance(value, bool):
        return "да" if value else "нет"
    return value


def _comment(cell, text: str, author: str = "compsheet") -> None:
    comment = Comment(text, author)
    comment.width = 320
    comment.height = 160
    cell.comment = comment


class BlockWriter:
    """Emits one block of ``suppliers_per_block`` offers at a given row."""

    def __init__(self, sheet, block: BlockTemplate, config: Config, fx: FxTable,
                 fx_reference: str):
        self.sheet = sheet
        self.block = block
        self.config = config
        self.fx = fx
        self.fx_reference = fx_reference

    # -- helpers --------------------------------------------------------

    def _value_column(self, supplier_index: int) -> int:
        return self.block.supplier_columns[supplier_index][0]

    def _emit_row(self, spec, target_row: int) -> None:
        if spec.height is not None:
            self.sheet.row_dimensions[target_row].height = spec.height
        for cell_spec in spec.cells:
            cell = self.sheet.cell(row=target_row, column=cell_spec.column)
            cell_spec.apply(cell)
            # Static text (labels, table headers) is part of the template; the
            # per-offer values are written afterwards and overwrite these.
            if cell_spec.value is not None:
                cell.value = cell_spec.value
            if cell_spec.merge_to:
                self.sheet.merge_cells(
                    start_row=target_row, start_column=cell_spec.column,
                    end_row=target_row, end_column=cell_spec.merge_to,
                )

    def _write(self, row: int, column: int, value, number_format: str | None = None):
        cell = self.sheet.cell(row=row, column=column)
        cell.value = value
        if number_format:
            cell.number_format = number_format
        return cell

    # -- the block ------------------------------------------------------

    def write(self, top: int, offers: list[Offer], item_count: int,
              placements: list | None = None) -> int:
        """Write one block; returns the first row after it."""
        block = self.block
        rows = block.rows
        item_index = block.item_row_index

        # Row plan: every template row once, with the nomenclature row repeated.
        plan: list[tuple[object, int | None]] = []
        for index, spec in enumerate(rows):
            if index == item_index:
                for position in range(item_count):
                    plan.append((spec, position))
            elif spec.kind == ITEM:
                continue  # extra sample rows are replaced by the repeat above
            else:
                plan.append((spec, None))

        target_rows: dict[int, int] = {}
        item_rows: list[int] = []
        row_number = top
        for spec, position in plan:
            self._emit_row(spec, row_number)
            if position is None:
                target_rows[rows.index(spec)] = row_number
            else:
                item_rows.append(row_number)
            row_number += 1
        bottom = row_number - 1

        role_rows = {
            spec.role: target_rows[index]
            for index, spec in enumerate(rows)
            if spec.role and index in target_rows
        }

        for index, spec in enumerate(rows):
            if spec.kind == ITEM:
                continue
            row = target_rows.get(index)
            if row is None:
                continue
            if spec.kind == HEADER:
                self._write_header(row, offers)
            elif spec.kind == PARAM:
                self._write_param(row, spec.role, offers)
            elif spec.kind == TOTAL:
                self._write_total(row, spec.role, offers, item_rows, role_rows)
            elif spec.kind in (ITEM_HEADER, OTHER):
                continue

        for position, row in enumerate(item_rows):
            self._write_item(row, position, offers)

        if placements is not None:
            for index, offer in enumerate(offers):
                placements.append(Placement(
                    company=offer.display_name,
                    sheet=self.sheet.title,
                    column=self._value_column(index),
                    role_rows=dict(role_rows),
                    item_rows=list(item_rows),
                ))

        return bottom + 1

    def _write_header(self, row: int, offers: list[Offer]) -> None:
        for index in range(self.block.suppliers_per_block):
            column = self._value_column(index)
            offer = offers[index] if index < len(offers) else None
            if offer is None:
                self._write(row, column, "—")
                continue
            cell = self._write(row, column, offer.display_name)
            notes = []
            if offer.source_files:
                notes.append("Файлы: " + ", ".join(
                    path.rsplit("/", 1)[-1] for path in offer.source_files))
            if offer.assumptions:
                notes.append("Допущения системы:\n- " + "\n- ".join(offer.assumptions))
            if offer.warnings:
                notes.append("Предупреждения:\n- " + "\n- ".join(offer.warnings))
            if offer.confidence is not None:
                notes.append(f"Уверенность извлечения: {offer.confidence:.0%}")
            if notes:
                _comment(cell, "\n\n".join(notes))

    def _write_param(self, row: int, role: str, offers: list[Offer]) -> None:
        for index in range(self.block.suppliers_per_block):
            column = self._value_column(index)
            offer = offers[index] if index < len(offers) else None
            if offer is None:
                self._write(row, column, None)
                continue
            field_value = getattr(offer, role, None)
            if field_value is None:
                continue
            value = field_value.value
            if role == "vat_included":
                value = (VAT_IN_PRICE if value is True
                         else VAT_NOT_IN_PRICE if value is False else None)
            cell = self._write(row, column, _field_text(value))
            if field_value.source in (ASSUMED, MISSING) and field_value.note:
                label = "ПРЕДПОЛОЖЕНО" if field_value.source == ASSUMED else "НЕ НАЙДЕНО"
                _comment(cell, f"{label}: {field_value.note}")

    def _write_item(self, row: int, position: int, offers: list[Offer]) -> None:
        shared = self.block.item_columns
        spec = self.config.spec[position] if position < len(self.config.spec) else None

        if "num" in shared:
            self._write(row, shared["num"], position + 1)
        if "description" in shared:
            self._write(row, shared["description"],
                        spec.description if spec else f"Позиция {position + 1}")
        if "unit" in shared:
            self._write(row, shared["unit"], spec.unit if spec else None)
        qty_column = shared["qty"]
        self._write(row, qty_column, spec.qty if spec else None)

        for index in range(self.block.suppliers_per_block):
            columns = self.block.supplier_item_columns[index]
            offer = offers[index] if index < len(offers) else None
            price_column = columns.get("unit_price_supplier") or columns.get("unit_price_buyer")
            amount_column = columns.get("amount")
            if offer is None:
                continue
            item = next((line for line in offer.line_items if line.spec_index == position), None)
            price = item.unit_price.value if item else None
            if price_column:
                cell = self._write(row, price_column,
                                   float(price) if price is not None else 0, "#,##0.00")
                if item is None:
                    _comment(cell, f"{self.config.not_offered_text}: позиция отсутствует в КП")
                elif item.unit_price.source in (ASSUMED, MISSING):
                    _comment(cell, f"{item.unit_price.source.upper()}: {item.unit_price.note}")
                if item is not None and item.description.value == self.config.not_offered_text:
                    _comment(cell, item.description.note or self.config.not_offered_text)
            if amount_column and price_column:
                self._write(
                    row, amount_column,
                    f"={_letter(qty_column)}{row}*{_letter(price_column)}{row}",
                    "#,##0.00",
                )

    def _write_total(self, row: int, role: str, offers: list[Offer],
                     item_rows: list[int], role_rows: dict[str, int]) -> None:
        first, last = item_rows[0], item_rows[-1]
        vat_rate = self.config.vat_rate

        for index in range(self.block.suppliers_per_block):
            column = self._value_column(index)
            letter = _letter(column)
            offer = offers[index] if index < len(offers) else None
            if offer is None:
                self._write(row, column, None)
                continue

            columns = self.block.supplier_item_columns[index]
            qty_letter = _letter(self.block.item_columns["qty"])
            amount_letter = _letter(columns.get("amount", column))
            currency_row = role_rows.get("currency")
            rate = (f"VLOOKUP({letter}{currency_row},{self.fx_reference},2,0)"
                    if currency_row else "1")

            if role == "total_qty":
                self._write(row, column, f"=SUM({qty_letter}{first}:{qty_letter}{last})")
            elif role == "total_amount_ccy":
                self._write(row, column,
                            f"=SUM({amount_letter}{first}:{amount_letter}{last})", "#,##0.00")
            elif role == "total_amount_base":
                source = role_rows.get("total_amount_ccy")
                if source:
                    self._write(row, column, f"={letter}{source}*{rate}", "#,##0.00")
            elif role == "vat_amount":
                base = role_rows.get("total_amount_base") or role_rows.get("total_amount_ccy")
                vat_row = role_rows.get("vat_included")
                if base and vat_row:
                    self._write(
                        row, column,
                        f'=IF({letter}{vat_row}="{VAT_NOT_IN_PRICE}",'
                        f'{letter}{base}*{vat_rate},0)',
                        "#,##0.00",
                    )
                elif base:
                    self._write(row, column, f"={letter}{base}*{vat_rate}", "#,##0.00")
            elif role == "transport_cost":
                value = offer.transport_cost.value
                if value:
                    self._write(row, column, f"={float(value)}*{rate}", "#,##0.00")
                else:
                    cell = self._write(row, column, 0, "#,##0.00")
                    if offer.transport_cost.source == MISSING:
                        _comment(cell, "НЕ НАЙДЕНО: транспортные расходы в КП не указаны, принят 0")
            elif role == "grand_total":
                parts = [
                    f"{letter}{role_rows[key]}"
                    for key in ("total_amount_base", "vat_amount", "transport_cost")
                    if key in role_rows
                ]
                if not parts and "total_amount_ccy" in role_rows:
                    parts = [f"{letter}{role_rows['total_amount_ccy']}"]
                self._write(row, column, "=" + "+".join(parts), "#,##0.00")
            else:
                field_value = getattr(offer, role, None)
                if field_value is not None:
                    cell = self._write(row, column, _field_text(field_value.value))
                    if field_value.source in (ASSUMED, MISSING) and field_value.note:
                        label = "ПРЕДПОЛОЖЕНО" if field_value.source == ASSUMED else "НЕ НАЙДЕНО"
                        _comment(cell, f"{label}: {field_value.note}")


def render(template_path: str, offers: list[Offer], config: Config, fx: FxTable,
           out_path: str) -> RenderReport:
    from openpyxl import load_workbook

    workbook = load_workbook(template_path)
    sheet, block = find_main_sheet(workbook, config)
    report = RenderReport(item_rows_per_block=max(1, len(config.spec)))

    for offer in offers:
        currency = offer.currency.value
        if not fx.has(str(currency or "")):
            raise FxError(
                f"{offer.display_name}: no exchange rate for {currency or '<not set>'}. "
                "Add it to the fx config or set the currency during review."
            )

    fx_reference, fx_rows = write_fx_table(workbook, fx, config)
    report.fx_rows = fx_rows

    validations = _capture_validations(sheet, block)
    _clear_from(sheet, block)

    per_block = config.suppliers_per_block or block.suppliers_per_block
    if per_block > block.suppliers_per_block:
        report.notes.append(
            f"шаблон рассчитан на {block.suppliers_per_block} поставщиков в блоке; "
            f"использую {block.suppliers_per_block} вместо {per_block}"
        )
        per_block = block.suppliers_per_block

    item_count = max(1, len(config.spec))
    writer = BlockWriter(sheet, block, config, fx, fx_reference)
    breaks: list[int] = []
    row = block.top
    for start in range(0, len(offers), per_block):
        group = offers[start:start + per_block]
        block_top = row
        row = writer.write(row, group, item_count, report.placements)
        for validation_type, formula, offsets in validations:
            validation = DataValidation(type=validation_type, formula1=formula, allow_blank=True)
            sheet.add_data_validation(validation)
            for top_offset, min_col, bottom_offset, max_col in offsets:
                validation.add(
                    f"{_letter(min_col)}{block_top + top_offset}:"
                    f"{_letter(max_col)}{block_top + bottom_offset}"
                )
        report.blocks += 1
        breaks.append(row - 1)
        row += 1  # spacer between blocks

    report.rows_written = row - block.top
    _finish_page_setup(sheet, block, breaks, config, report)
    workbook.save(out_path)
    return report


def _finish_page_setup(sheet, block: BlockTemplate, breaks: list[int], config: Config,
                       report: RenderReport) -> None:
    """Undo "fit everything on one page" and break pages between blocks."""
    last_row = max(breaks) if breaks else block.bottom
    sheet.print_area = f"A1:{_letter(block.max_column)}{last_row}"
    sheet.page_setup.fitToWidth = 1
    if sheet.page_setup.fitToHeight not in (0, None):
        report.notes.append(
            "в шаблоне была включена подгонка под одну страницу по высоте — "
            "сброшена, иначе лист сжимается до нечитаемого масштаба"
        )
    sheet.page_setup.fitToHeight = 0
    if sheet.sheet_properties.pageSetUpPr is not None:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{block.top - 1}" if block.top > 1 else None
    if config.page_break_between_blocks:
        for row in breaks[:-1]:
            sheet.row_breaks.append(Break(id=row + 1))
