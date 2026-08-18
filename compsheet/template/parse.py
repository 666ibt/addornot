"""Read the geometry of a comparison block out of the template workbook.

Nothing about the block is hardcoded: the label column, the number of parameter
rows, how many columns belong to one supplier, where the nomenclature table
starts and which totals follow it are all discovered by matching the template's
own labels against the roles declared in the config.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from ..config import PARAM_ROLES, TOTAL_ROLES, Config
from ..errors import TemplateError

# Row kinds, in the order they appear inside a block.
HEADER, PARAM, ITEM_HEADER, ITEM, TOTAL, OTHER, SPACER = (
    "header", "param", "item_header", "item", "total", "other", "spacer"
)


@dataclass
class CellSpec:
    """One cell of the block template: its style, its static text, its merge."""

    column: int
    value: Any = None
    style: Any = None
    merge_to: int | None = None

    def apply(self, cell) -> None:
        if self.style is not None:
            cell._style = copy(self.style)


@dataclass
class RowSpec:
    kind: str
    role: str | None = None
    height: float | None = None
    cells: list[CellSpec] = field(default_factory=list)

    def cell(self, column: int) -> CellSpec | None:
        return next((c for c in self.cells if c.column == column), None)


@dataclass
class BlockTemplate:
    sheet_title: str
    top: int
    bottom: int
    rows: list[RowSpec]
    label_column: int
    supplier_columns: list[tuple[int, int]]
    item_columns: dict[str, int]
    supplier_item_columns: list[dict[str, int]]
    max_column: int

    @property
    def suppliers_per_block(self) -> int:
        return len(self.supplier_columns)

    @property
    def item_row_index(self) -> int:
        for index, row in enumerate(self.rows):
            if row.kind == ITEM:
                return index
        raise TemplateError("block has no nomenclature row")

    def rows_of(self, kind: str) -> list[RowSpec]:
        return [row for row in self.rows if row.kind == kind]

    def role_row(self, role: str) -> RowSpec | None:
        return next((row for row in self.rows if row.role == role), None)


def _merged_map(sheet: Worksheet) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    """(row, col) of a merge anchor -> its bounds."""
    anchors = {}
    for merged in sheet.merged_cells.ranges:
        bounds = merged.min_row, merged.min_col, merged.max_row, merged.max_col
        anchors[(merged.min_row, merged.min_col)] = bounds
    return anchors


def _find_label_column(sheet: Worksheet, config: Config) -> tuple[int, int]:
    """Locate the column holding role labels and the first parameter row."""
    for row in range(1, min(sheet.max_row, 300) + 1):
        for column in range(1, min(sheet.max_column, 20) + 1):
            value = sheet.cell(row=row, column=column).value
            if isinstance(value, str) and config.match_role(value, PARAM_ROLES):
                return column, row
    raise TemplateError(
        "could not find any parameter label from the config in the template; "
        "check label_patterns or the template layout"
    )


def _supplier_columns(sheet: Worksheet, row: int, label_column: int,
                      merges: dict) -> list[tuple[int, int]]:
    """Column spans belonging to each supplier, taken from the merge pattern."""
    label_bounds = merges.get((row, label_column))
    first_data = (label_bounds[3] if label_bounds else label_column) + 1

    spans: list[tuple[int, int]] = []
    column = first_data
    while column <= sheet.max_column:
        bounds = merges.get((row, column))
        if bounds:
            spans.append((bounds[1], bounds[3]))
            column = bounds[3] + 1
        else:
            spans.append((column, column))
            column += 1
    if not spans:
        raise TemplateError("no supplier data columns found to the right of the labels")

    # Uniform spans are the norm; a trailing odd one is a notes column, not a
    # supplier, so drop it.
    width = spans[0][1] - spans[0][0] + 1
    uniform = [span for span in spans if span[1] - span[0] + 1 == width]
    return uniform or spans


def parse_block(sheet: Worksheet, config: Config) -> BlockTemplate:
    merges = _merged_map(sheet)
    label_column, first_param = _find_label_column(sheet, config)
    supplier_columns = _supplier_columns(sheet, first_param, label_column, merges)
    max_column = supplier_columns[-1][1]

    # The company header sits directly above the first parameter row.
    top = first_param
    header_candidate = first_param - 1
    if header_candidate >= 1:
        anchor = supplier_columns[0][0]
        if sheet.cell(row=header_candidate, column=anchor).value is not None or \
                (header_candidate, anchor) in merges:
            top = header_candidate

    # Walk down: parameters, the nomenclature header, its rows, then totals.
    rows: list[RowSpec] = []
    item_columns: dict[str, int] = {}
    supplier_item_columns: list[dict[str, int]] = []
    row_number = top
    seen_item_header = False
    seen_total = False
    bottom = top

    while row_number <= sheet.max_row:
        label = sheet.cell(row=row_number, column=label_column).value
        param_role = config.match_role(label, PARAM_ROLES) if isinstance(label, str) else None
        total_role = config.match_role(label, TOTAL_ROLES) if isinstance(label, str) else None
        item_role_hits = [
            config.match_item_column(sheet.cell(row=row_number, column=column).value)
            for column in range(1, max_column + 1)
        ]
        is_item_header = (
            not seen_item_header
            and sum(1 for role in item_role_hits if role) >= 2
            and param_role is None
        )

        if is_item_header:
            kind, role = ITEM_HEADER, None
            seen_item_header = True
            for column, item_role in enumerate(item_role_hits, start=1):
                if not item_role:
                    continue
                inside = next((index for index, (start, end) in enumerate(supplier_columns)
                               if start <= column <= end), None)
                if inside is None:
                    item_columns.setdefault(item_role, column)
                else:
                    while len(supplier_item_columns) <= inside:
                        supplier_item_columns.append({})
                    supplier_item_columns[inside].setdefault(item_role, column)
        elif total_role:
            kind, role, seen_total = TOTAL, total_role, True
        elif param_role and not seen_item_header:
            kind, role = PARAM, param_role
        elif row_number == top and not seen_item_header:
            kind, role = HEADER, None
        elif seen_item_header and not seen_total:
            kind, role = ITEM, None
        elif seen_total:
            # After the totals, an unlabelled row ends the block.
            if label is None and not any(
                sheet.cell(row=row_number, column=column).value is not None
                for column in range(1, max_column + 1)
            ):
                break
            kind, role = OTHER, None
        else:
            kind, role = OTHER, None

        rows.append(_capture_row(sheet, row_number, kind, role, max_column, merges))
        bottom = row_number
        row_number += 1

    if not any(row.kind == ITEM for row in rows):
        raise TemplateError("no nomenclature rows found between the table header and the totals")
    if not supplier_item_columns or not any(supplier_item_columns):
        raise TemplateError(
            "could not identify the price/amount columns of a supplier in the table header"
        )
    for shared in ("qty",):
        if shared not in item_columns:
            raise TemplateError(f"the nomenclature table has no {shared!r} column")

    return BlockTemplate(
        sheet_title=sheet.title,
        top=top,
        bottom=bottom,
        rows=rows,
        label_column=label_column,
        supplier_columns=supplier_columns,
        item_columns=item_columns,
        supplier_item_columns=supplier_item_columns,
        max_column=max_column,
    )


def _capture_row(sheet: Worksheet, row_number: int, kind: str, role: str | None,
                 max_column: int, merges: dict) -> RowSpec:
    dimension = sheet.row_dimensions.get(row_number)
    spec = RowSpec(
        kind=kind,
        role=role,
        height=dimension.height if dimension is not None else None,
    )
    for column in range(1, max_column + 1):
        cell = sheet.cell(row=row_number, column=column)
        bounds = merges.get((row_number, column))
        value = cell.value
        # Formulas are regenerated from the row's role, never copied.
        if isinstance(value, str) and value.startswith("="):
            value = None
        spec.cells.append(
            CellSpec(
                column=column,
                value=value,
                style=copy(cell._style),
                merge_to=bounds[3] if bounds and bounds[3] > column else None,
            )
        )
    return spec
