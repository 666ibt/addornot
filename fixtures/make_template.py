#!/usr/bin/env python3
"""Generate a template workbook shaped like the hand-made ones.

Deliberately reproduces the quirks the pipeline has to cope with:

* one sample comparison block for two suppliers, with merged label cells;
* a single nomenclature row (real purchases have several);
* totals that SUM exactly that one row instead of a range;
* a local-currency conversion that VLOOKUPs an *empty* rate table, so the
  formula evaluates to #N/A in the original file;
* page setup pinned to "fit to one page", which crushes a multi-block sheet.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MAIN = "Конкурентный лист"
REFERENCE = "Справочники"

THIN = Side(style="thin", color="FF808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="FFD9E1F2")
LABEL_FILL = PatternFill("solid", fgColor="FFF2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="FFFFF2CC")
BOLD = Font(name="Calibri", size=10, bold=True)
NORMAL = Font(name="Calibri", size=10)
WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONEY = "#,##0.00"

PARAMETERS = [
    "Принадлежность поставщика",
    "Производитель (да/нет)",
    "Техническое заключение",
    "Страна происхождения",
    "Страна отгрузки",
    "Город (пункт отправления)",
    "Валюта ТКП",
    "Дата предложения",
    "Срок действия предложения",
    "Базис поставки (ИНКОТЕРМС 2020)",
    "Пункт поставки",
    "НДС в цене",
]

TOTALS = [
    ("Итого количество", "qty"),
    ("Сумма в валюте ТКП", "money"),
    ("Сумма в UZS", "money"),
    ("НДС 12%", "money"),
    ("Транспортные расходы", "money"),
    ("ИТОГО с НДС и транспортом", "money"),
    ("Срок поставки", "text"),
    ("Условия оплаты", "text"),
    ("Гарантийные обязательства", "text"),
    ("Процедура закупки", "text"),
]

COUNTRIES = ["Узбекистан", "Россия", "Казахстан", "Турция", "Китай", "Германия", "Италия", "Индия"]
CURRENCIES = ["UZS", "USD", "EUR", "RUB", "TRY", "CNY"]
INCOTERMS = ["EXW", "FCA", "FAS", "FOB", "CFR", "CIF", "CPT", "CIP", "DAP", "DPU", "DDP"]
UNITS = ["шт", "компл", "м", "м2", "м3", "кг", "т", "л"]


def build_reference(sheet):
    columns = [("Страны", COUNTRIES), ("Валюты", CURRENCIES),
               ("Базис поставки", INCOTERMS), ("Единицы измерения", UNITS)]
    for index, (title, values) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        for offset, value in enumerate(values, start=2):
            sheet.cell(row=offset, column=index, value=value)
        sheet.column_dimensions[get_column_letter(index)].width = 20

    # The rate table the main sheet points at -- header only, no rows.
    sheet["F1"] = "Валюта"
    sheet["G1"] = "Курс к UZS"
    sheet["F1"].font = sheet["G1"].font = BOLD
    sheet["F1"].fill = sheet["G1"].fill = HEADER_FILL
    sheet.column_dimensions["F"].width = 14
    sheet.column_dimensions["G"].width = 16


def build_block(sheet, top: int, suppliers: int = 2) -> int:
    """Write one comparison block starting at ``top``; return the next free row."""
    first_data_column = 5  # E
    columns_per_supplier = 2

    def supplier_columns(index: int) -> tuple[int, int]:
        start = first_data_column + index * columns_per_supplier
        return start, start + columns_per_supplier - 1

    row = top
    # -- company header ---------------------------------------------------
    sheet.cell(row=row, column=2, value="Поставщик").font = BOLD
    sheet.cell(row=row, column=2).fill = HEADER_FILL
    sheet.cell(row=row, column=2).alignment = CENTER
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    for index in range(suppliers):
        start, end = supplier_columns(index)
        cell = sheet.cell(row=row, column=start, value=f"Поставщик {index + 1}")
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    for column in range(1, supplier_columns(suppliers - 1)[1] + 1):
        sheet.cell(row=row, column=column).border = BORDER
    sheet.row_dimensions[row].height = 30
    row += 1

    # -- parameter rows ---------------------------------------------------
    for label in PARAMETERS:
        cell = sheet.cell(row=row, column=1, value=label)
        cell.font = NORMAL
        cell.fill = LABEL_FILL
        cell.alignment = WRAP
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        for index in range(suppliers):
            start, end = supplier_columns(index)
            value = sheet.cell(row=row, column=start)
            value.font = NORMAL
            value.alignment = WRAP
            sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        for column in range(1, supplier_columns(suppliers - 1)[1] + 1):
            sheet.cell(row=row, column=column).border = BORDER
        sheet.row_dimensions[row].height = 22
        row += 1

    # -- nomenclature table header ---------------------------------------
    headers = ["№", "Наименование", "Ед. изм.", "Кол-во"]
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font, cell.fill, cell.alignment, cell.border = BOLD, HEADER_FILL, CENTER, BORDER
    for index in range(suppliers):
        start, end = supplier_columns(index)
        for column, title in zip((start, end), ("Цена у поставщика", "Сумма")):
            cell = sheet.cell(row=row, column=column, value=title)
            cell.font, cell.fill, cell.alignment, cell.border = BOLD, HEADER_FILL, CENTER, BORDER
    sheet.row_dimensions[row].height = 32
    row += 1

    # -- a single nomenclature row (real sheets need several) -------------
    item_row = row
    sheet.cell(row=row, column=1, value=1).alignment = CENTER
    sheet.cell(row=row, column=2, value="Наименование позиции").alignment = WRAP
    sheet.cell(row=row, column=3, value="шт").alignment = CENTER
    sheet.cell(row=row, column=4, value=1).alignment = CENTER
    for index in range(suppliers):
        start, end = supplier_columns(index)
        price = sheet.cell(row=row, column=start)
        price.number_format = MONEY
        amount = sheet.cell(row=row, column=end)
        amount.value = f"={get_column_letter(4)}{row}*{get_column_letter(start)}{row}"
        amount.number_format = MONEY
    for column in range(1, supplier_columns(suppliers - 1)[1] + 1):
        cell = sheet.cell(row=row, column=column)
        cell.font = NORMAL
        cell.border = BORDER
    sheet.row_dimensions[row].height = 30
    row += 1

    # -- totals -----------------------------------------------------------
    for label, kind in TOTALS:
        cell = sheet.cell(row=row, column=1, value=label)
        cell.font = BOLD
        cell.fill = TOTAL_FILL
        cell.alignment = WRAP
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        for index in range(suppliers):
            start, end = supplier_columns(index)
            value = sheet.cell(row=row, column=start)
            value.font = BOLD
            value.alignment = CENTER if kind != "text" else WRAP
            if kind == "money":
                value.number_format = MONEY
            sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
            amount_letter = get_column_letter(end)
            if label.startswith("Итого количество"):
                value.value = f"=D{item_row}"
            elif label.startswith("Сумма в валюте"):
                value.value = f"={amount_letter}{item_row}"
            elif label.startswith("Сумма в UZS"):
                # Points at an empty rate table: #N/A in the original file.
                value.value = (
                    f"={get_column_letter(start)}{row - 1}"
                    f"*VLOOKUP({get_column_letter(start)}{top + 7},"
                    f"'{REFERENCE}'!$F$2:$G$2,2,0)"
                )
            elif label.startswith("НДС"):
                value.value = f"={get_column_letter(start)}{row - 1}*0.12"
            elif label.startswith("ИТОГО"):
                value.value = (
                    f"={get_column_letter(start)}{row - 3}+{get_column_letter(start)}{row - 2}"
                    f"+{get_column_letter(start)}{row - 1}"
                )
        for column in range(1, supplier_columns(suppliers - 1)[1] + 1):
            sheet.cell(row=row, column=column).border = BORDER
        sheet.row_dimensions[row].height = 22
        row += 1

    return row + 1  # one spacer row between blocks


def main(out: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MAIN
    reference = workbook.create_sheet(REFERENCE)
    build_reference(reference)

    sheet["A1"] = "КОНКУРЕНТНЫЙ ЛИСТ"
    sheet["A1"].font = Font(name="Calibri", size=14, bold=True)
    sheet["A2"] = "Предмет закупки:"
    sheet["A2"].font = BOLD

    widths = {"A": 6, "B": 46, "C": 9, "D": 10, "E": 18, "F": 18, "G": 18, "H": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    next_row = build_block(sheet, top=4)

    # Dropdowns sourced from the reference sheet.
    validations = [
        (f"'{REFERENCE}'!$A$2:$A$9", ["E8:H9"]),
        (f"'{REFERENCE}'!$B$2:$B$7", ["E11:H11"]),
        (f"'{REFERENCE}'!$C$2:$C$12", ["E14:H14"]),
    ]
    for formula, ranges in validations:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(validation)
        for cells in ranges:
            validation.add(cells)

    sheet.freeze_panes = "E4"
    sheet.print_area = f"A1:H{next_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1  # crushes the sheet once blocks are added
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures/template.xlsx"))
