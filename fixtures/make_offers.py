#!/usr/bin/env python3
"""Generate a regression fixture set of supplier offers.

The set deliberately covers every awkward case the pipeline claims to handle:
a clean text PDF, a scan with no text layer, a readable .docx, a .docx whose
text extraction comes out as legacy-encoding soup, a supplier split across two
files, a supplier that quoted only one of the two required positions, a
supplier that priced a single unit instead of the required quantity, and a
truncated download.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

FONT = "DejaVu Sans"


def soffice(args: list[str], outdir: Path):
    profile = (outdir / "_loprofile").resolve()
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--nolockcheck",
         f"-env:UserInstallation=file://{profile}", *args,
         "--outdir", str(outdir)],
        check=True, capture_output=True, timeout=300,
    )


def html_to_pdf(html: str, out: Path) -> Path:
    src = out.with_suffix(".html")
    src.write_text(html, encoding="utf-8")
    soffice(["--convert-to", "pdf", str(src)], out.parent)
    src.unlink()
    produced = out.parent / (out.stem + ".pdf")
    return produced


def page(title: str, body_rows: list[str], footer: list[str]) -> str:
    rows = "\n".join(body_rows)
    tail = "".join(f"<p>{line}</p>" for line in footer)
    return f"""<html><head><meta charset="utf-8"></head>
<body style="font-family:'{FONT}';font-size:11pt">
<h2>{title}</h2>
<table border="1" cellspacing="0" cellpadding="4" width="100%">
<tr><th>№</th><th>Наименование</th><th>Ед.</th><th>Кол-во</th><th>Цена за ед.</th><th>Сумма</th></tr>
{rows}
</table>
{tail}
</body></html>"""


def item_row(n, name, unit, qty, price, total):
    return (f"<tr><td>{n}</td><td>{name}</td><td>{unit}</td><td>{qty}</td>"
            f"<td>{price}</td><td>{total}</td></tr>")


def make_alfa(outdir: Path):
    """Clean text PDF, USD, prices without VAT, both positions."""
    html = page(
        "Коммерческое предложение № 114 от 12.03.2025<br>ООО «Альфа Нефтемаш», г. Ташкент, Узбекистан",
        [
            item_row(1, "Резервуар горизонтальный стальной РГС-50, 50 м3", "шт", "4", "18 500,00", "74 000,00"),
            item_row(2, "Резервуар вертикальный стальной РВС-100, 100 м3", "шт", "2", "31 200,00", "62 400,00"),
        ],
        [
            "Итого: 136 400,00 USD",
            "Цены указаны без НДС. Валюта предложения — доллар США (USD).",
            "Базис поставки: DDP г. Ташкент (Инкотермс 2020).",
            "Страна происхождения товара: Узбекистан. Страна отгрузки: Узбекистан.",
            "Срок поставки: 45 календарных дней с даты предоплаты.",
            "Условия оплаты: 50% предоплата, 50% по факту поставки.",
            "Гарантия: 24 месяца с даты отгрузки.",
            "Предложение действительно до 12.05.2025.",
            "Мы являемся производителем предлагаемого оборудования.",
        ],
    )
    html_to_pdf(html, outdir / "alfa_neftemash_kp")


def make_beta_scan(outdir: Path):
    """Same content rendered to images: a scan with no text layer."""
    from PIL import Image

    html = page(
        "ТИЖОРАТ ТАКЛИФИ / Коммерческое предложение № 7-Р от 05.03.2025<br>ООО «Бета Резервуар», г. Самарканд, Узбекистан",
        [
            item_row(1, "Резервуар горизонтальный стальной РГС-50, 50 м3", "шт", "4", "235 000 000", "940 000 000"),
            item_row(2, "Резервуар вертикальный стальной РВС-100, 100 м3", "шт", "2", "398 000 000", "796 000 000"),
        ],
        [
            "Итого с НДС: 1 736 000 000 сум",
            "Валюта: UZS (сум). Цены указаны с учетом НДС 12%.",
            "Базис поставки: FCA г. Самарканд.",
            "Страна происхождения: Узбекистан. Отгрузка: Узбекистан.",
            "Срок поставки: 60 дней. Оплата: 100% предоплата.",
            "Гарантия: 12 месяцев.",
        ],
    )
    pdf = html_to_pdf(html, outdir / "_beta_tmp")
    subprocess.run(["pdftoppm", "-r", "150", "-png", str(pdf), str(outdir / "_beta_page")],
                   check=True, capture_output=True)
    pages = sorted(outdir.glob("_beta_page-*.png"))
    images = [Image.open(p).convert("RGB") for p in pages]
    images[0].save(outdir / "beta_rezervuar_skan.pdf", save_all=True, append_images=images[1:])
    for path in [pdf, *pages]:
        path.unlink()


def make_gamma(outdir: Path):
    """Readable .docx in Uzbek/Latin, EUR, only one of the two positions."""
    import docx

    document = docx.Document()
    document.add_paragraph("GAMMA SAVDO MChJ, Toshkent shahri, O'zbekiston")
    document.add_paragraph("Tijorat taklifi № 42, sana: 18.03.2025")
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Nomi", "Birlik", "Miqdor", "Narx", "Summa"]):
        cell.text = text
    row = table.add_row().cells
    for cell, text in zip(row, ["Rezervuar gorizontal po'lat RGS-50, 50 m3", "dona", "4", "16 900,00", "67 600,00"]):
        cell.text = text
    for line in [
        "Jami: 67 600,00 EUR",
        "Valyuta: EUR. Narxlar QQS (NDS) siz ko'rsatilgan.",
        "Yetkazib berish sharti: CPT Toshkent.",
        "Kelib chiqish mamlakati: Turkiya. Jo'natish mamlakati: Turkiya.",
        "Yetkazib berish muddati: 90 kun. To'lov: 30% oldindan.",
        "Kafolat: 18 oy.",
    ]:
        document.add_paragraph(line)
    document.save(outdir / "gamma_savdo_taklif.docx")


def make_delta_mangled(outdir: Path):
    """A .docx whose text extraction comes out as legacy-encoding soup.

    The supplier also priced a *single* unit while the specification asks for
    four — the pipeline must flag that instead of comparing the totals.
    """
    import docx

    def soup(text: str) -> str:
        return text.encode("cp1251").decode("latin-1")

    document = docx.Document()
    for line in [
        "ООО «Дельта Металл», г. Москва, Россия",
        "Коммерческое предложение № 512 от 01.03.2025",
        "Резервуар горизонтальный стальной РГС-50, 50 м3 — 1 шт — 1 750 000,00 RUB",
        "Резервуар вертикальный стальной РВС-100, 100 м3 — 1 шт — 2 980 000,00 RUB",
        "Итого: 4 730 000,00 RUB",
        "Валюта: RUB. Цены с НДС 20%.",
        "Базис поставки: EXW г. Москва. Страна происхождения: Россия.",
        "Срок поставки: 30 дней. Оплата: по факту. Гарантия: 12 месяцев.",
    ]:
        document.add_paragraph(soup(line))
    document.save(outdir / "delta_metall_kp.docx")

    # The same supplier also sent a signed scan; the readable copy of the data
    # therefore only exists behind OCR.
    from PIL import Image

    html = page(
        "ООО «Дельта Металл», г. Москва, Россия<br>Коммерческое предложение № 512 от 01.03.2025",
        [
            item_row(1, "Резервуар горизонтальный стальной РГС-50, 50 м3", "шт", "1", "1 750 000,00", "1 750 000,00"),
            item_row(2, "Резервуар вертикальный стальной РВС-100, 100 м3", "шт", "1", "2 980 000,00", "2 980 000,00"),
        ],
        [
            "Итого: 4 730 000,00 RUB",
            "Валюта: RUB. Цены указаны с НДС 20%.",
            "Базис поставки: EXW г. Москва. Страна происхождения: Россия.",
            "Срок поставки: 30 дней. Оплата: по факту поставки. Гарантия: 12 месяцев.",
        ],
    )
    pdf = html_to_pdf(html, outdir / "_delta_tmp")
    subprocess.run(["pdftoppm", "-r", "150", "-png", str(pdf), str(outdir / "_delta_page")],
                   check=True, capture_output=True)
    pages = sorted(outdir.glob("_delta_page-*.png"))
    images = [Image.open(p).convert("RGB") for p in pages]
    images[0].save(outdir / "delta_metall_skan.pdf", save_all=True, append_images=images[1:])
    for path in [pdf, *pages]:
        path.unlink()


def make_epsilon_two_files(outdir: Path):
    """One supplier, two files: a technical part and a commercial part."""
    tech = f"""<html><head><meta charset="utf-8"></head>
<body style="font-family:'{FONT}';font-size:11pt">
<h2>ООО «Эпсилон Инжиниринг» — Техническая часть предложения</h2>
<p>Предлагаемое оборудование: резервуары стальные по ГОСТ 17032-2010.</p>
<p>Изготовитель: ООО «Эпсилон Инжиниринг», г. Бухара, Узбекистан.</p>
<p>Техническое заключение заказчика: получено, положительное.</p>
<p>Гарантийный срок: 36 месяцев с даты ввода в эксплуатацию.</p>
</body></html>"""
    html_to_pdf(tech, outdir / "epsilon_tech")

    commercial = page(
        "ООО «Эпсилон Инжиниринг» — Коммерческая часть, КП № 88 от 20.03.2025",
        [
            item_row(1, "Резервуар горизонтальный стальной РГС-50, 50 м3", "шт", "4", "19 900,00", "79 600,00"),
            item_row(2, "Резервуар вертикальный стальной РВС-100, 100 м3", "шт", "2", "29 800,00", "59 600,00"),
        ],
        [
            "Итого: 139 200,00 USD без НДС.",
            "Базис поставки: DAP г. Ташкент. Страна отгрузки: Узбекистан.",
            "Срок поставки: 50 дней. Оплата: 100% по факту поставки.",
            "Транспортные расходы: включены в стоимость.",
        ],
    )
    html_to_pdf(commercial, outdir / "epsilon_commercial")


def make_broken(outdir: Path):
    """A truncated download that must be reported, not silently skipped."""
    source = outdir / "alfa_neftemash_kp.pdf"
    data = source.read_bytes()
    (outdir / "zeta_broken_download.pdf").write_bytes(data[: len(data) // 3])


def main(outdir: Path):
    if outdir.exists():
        shutil.rmtree(outdir)
    outdir.mkdir(parents=True)
    make_alfa(outdir)
    make_beta_scan(outdir)
    make_gamma(outdir)
    make_delta_mangled(outdir)
    make_epsilon_two_files(outdir)
    make_broken(outdir)
    for path in sorted(outdir.iterdir()):
        if path.is_dir():
            shutil.rmtree(path)
        else:
            print(f"{path.name:34s} {path.stat().st_size:>9,} bytes")


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures/offers"))
